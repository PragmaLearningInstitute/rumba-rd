from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Callable, List, Tuple
from xml.sax.saxutils import escape as xml_escape

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox
    from tkinter import font as tkfont
except ModuleNotFoundError as exc:
    if exc.name == "_tkinter":
        print(
            "\nTkinter n'est pas installé pour ce Python.\n"
            "Sur macOS Homebrew, exécutez:\n"
            "  brew install python-tk@3.14\n"
            "Puis relancez:\n"
            "  python3 app.py\n"
        )
        raise SystemExit(1)
    raise

from Function import FormatResult, gaussReformulator
from formatter import OutputStyle, render_html_document

APP_TITLE = "Rumba Reader"
BG = "#F2F5FB"
CARD = "#FFFFFF"
TEXT = "#151E35"
MUTED = "#667A99"
BORDER = "#C9D3E3"
INPUT_DARK = "#1F1F1F"
INPUT_DARK_TEXT = "#FFFFFF"
INPUT_LIGHT = "#FFFFFF"
INPUT_LIGHT_TEXT = "#6B7280"

FONT_OPTIONS = ["Avenir Next", "Arial", "OpenDyslexic"]
SIZE_OPTIONS = [11, 12, 13, 14, 15]


def get_base_dir() -> Path:
    if hasattr(sys, "_MEIPASS"):
        return Path(getattr(sys, "_MEIPASS"))
    return Path(__file__).resolve().parent


def install_packaged_fonts() -> None:
    fonts_dir = get_base_dir() / "assets" / "fonts"
    if not fonts_dir.exists():
        return

    files = [p for p in fonts_dir.iterdir() if p.suffix.lower() in {".ttf", ".otf", ".ttc"}]
    if not files:
        return

    try:
        if sys.platform == "darwin":
            target = Path.home() / "Library" / "Fonts"
            target.mkdir(parents=True, exist_ok=True)
            for src in files:
                dst = target / src.name
                if not dst.exists():
                    dst.write_bytes(src.read_bytes())
        elif os.name == "nt":
            local = os.environ.get("LOCALAPPDATA")
            if local:
                target = Path(local) / "Microsoft" / "Windows" / "Fonts"
                target.mkdir(parents=True, exist_ok=True)
                for src in files:
                    dst = target / src.name
                    if not dst.exists():
                        dst.write_bytes(src.read_bytes())
    except Exception:
        pass


def rounded_rect(canvas: tk.Canvas, x1: int, y1: int, x2: int, y2: int, radius: int, **kwargs) -> int:
    points = [
        x1 + radius,
        y1,
        x2 - radius,
        y1,
        x2,
        y1,
        x2,
        y1 + radius,
        x2,
        y2 - radius,
        x2,
        y2,
        x2 - radius,
        y2,
        x1 + radius,
        y2,
        x1,
        y2,
        x1,
        y2 - radius,
        x1,
        y1 + radius,
        x1,
        y1,
    ]
    return canvas.create_polygon(points, smooth=True, **kwargs)


def parse_markdown_segments(markdown: str) -> List[Tuple[str, bool]]:
    segs: List[Tuple[str, bool]] = []
    buf: List[str] = []
    bold = False
    i = 0
    while i < len(markdown):
        if markdown[i : i + 2] == "**":
            if buf:
                segs.append(("".join(buf), bold))
                buf = []
            bold = not bold
            i += 2
        else:
            buf.append(markdown[i])
            i += 1
    if buf:
        segs.append(("".join(buf), bold))
    return segs


def split_segments_by_line(segments: List[Tuple[str, bool]]) -> List[List[Tuple[str, bool]]]:
    lines: List[List[Tuple[str, bool]]] = [[]]
    for text, is_bold in segments:
        parts = text.split("\n")
        for idx, part in enumerate(parts):
            if part:
                lines[-1].append((part, is_bold))
            if idx < len(parts) - 1:
                lines.append([])
    return lines


class RoundSelect(tk.Canvas):
    def __init__(self, parent: tk.Widget, values: List[str], value: str, width: int, on_change: Callable[[str], None]) -> None:
        super().__init__(parent, width=width, height=42, bg=BG, highlightthickness=0, bd=0, cursor="arrow")
        self.values = values
        self.value = value
        self.on_change = on_change
        self.width_px = width
        self.menu = tk.Menu(self, tearoff=0)
        for v in values:
            self.menu.add_command(label=v, command=lambda vv=v: self.set_value(vv))
        self.bind("<Button-1>", self._open_menu)
        self._draw()

    def _draw(self) -> None:
        self.delete("all")
        rounded_rect(self, 1, 1, self.width_px - 1, 41, 12, fill="#FFFFFF", outline=BORDER, width=2)
        self.create_text(12, 21, text=self.value, anchor="w", fill=TEXT, font=("Arial", 10, "bold"))
        self.create_text(self.width_px - 14, 21, text="▾", anchor="e", fill=MUTED, font=("Arial", 12, "bold"))

    def _open_menu(self, _event) -> None:
        x = self.winfo_rootx()
        y = self.winfo_rooty() + self.winfo_height() + 2
        self.menu.tk_popup(x, y)

    def set_value(self, value: str) -> None:
        self.value = value
        self._draw()
        self.on_change(value)


class TextButton(tk.Canvas):
    def __init__(self, parent: tk.Widget, text: str, width: int, command: Callable[[], None], accent: bool = False) -> None:
        super().__init__(parent, width=width, height=42, bg=BG, highlightthickness=0, bd=0, cursor="arrow")
        self.text = text
        self.width_px = width
        self.command = command
        self.accent = accent
        self.hover = False
        self.bind("<Button-1>", self._click)
        self.bind("<Enter>", self._enter)
        self.bind("<Leave>", self._leave)
        self._draw()

    def set_text(self, text: str) -> None:
        self.text = text
        self._draw()

    def _draw(self) -> None:
        self.delete("all")
        if self.accent:
            fill = "#EAF2FF" if not self.hover else "#DBE8FF"
            outline = "#8EA6D1"
            color = "#102448" if not self.hover else "#1A2F5B"
        else:
            fill = "#FFFFFF"
            outline = BORDER
            color = "#0E162B" if not self.hover else "#394B73"
        rounded_rect(self, 1, 1, self.width_px - 1, 41, 12, fill=fill, outline=outline, width=2)
        self.create_text(self.width_px // 2, 21, text=self.text, fill=color, font=("Arial", 10, "bold"))

    def _click(self, _event) -> None:
        self.command()

    def _enter(self, _event) -> None:
        self.hover = True
        self._draw()

    def _leave(self, _event) -> None:
        self.hover = False
        self._draw()


class RoundedCard(tk.Canvas):
    def __init__(self, parent: tk.Widget, radius: int = 14, pad: int = 12) -> None:
        super().__init__(parent, bg=BG, highlightthickness=0, bd=0)
        self.radius = radius
        self.pad = pad
        self.inner = tk.Frame(self, bg=CARD)
        self.window_id = self.create_window(self.pad, self.pad, anchor="nw", window=self.inner)
        self.bind("<Configure>", self._on_configure)

    def _on_configure(self, _event) -> None:
        self.redraw()

    def redraw(self) -> None:
        self.delete("shape")
        w = self.winfo_width()
        h = self.winfo_height()
        if w < 20 or h < 20:
            return
        rounded_rect(self, 1, 1, w - 1, h - 1, self.radius, fill=CARD, outline=BORDER, width=2, tags="shape")
        self.coords(self.window_id, self.pad, self.pad)
        self.itemconfigure(self.window_id, width=max(10, w - self.pad * 2), height=max(10, h - self.pad * 2))


class RumbaReaderApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1600x900")
        self.root.minsize(1100, 700)
        self.root.configure(bg=BG, cursor="arrow")

        install_packaged_fonts()

        self.font_name = FONT_OPTIONS[1]
        self.font_size = 13
        self.input_dark_mode = True
        self.last_result: FormatResult | None = None

        self.font_label: tk.Label | None = None
        self.download_menu: tk.Menu | None = None

        self._build_ui()
        self._apply_input_theme()

    def _build_ui(self) -> None:
        main = tk.Frame(self.root, bg=BG)
        main.pack(fill="both", expand=True, padx=8, pady=8)

        tk.Label(main, text="Rumba Reader", bg=BG, fg=TEXT, font=("Arial", 38, "bold")).pack(anchor="w", pady=(0, 8))

        controls_line = tk.Frame(main, bg=BG)
        controls_line.pack(fill="x", pady=(0, 8))

        tk.Label(controls_line, text="Police", bg=BG, fg=TEXT, font=("Arial", 15, "bold")).pack(side="left", padx=(0, 8))
        self.font_select = RoundSelect(controls_line, FONT_OPTIONS, self.font_name, width=220, on_change=self._on_font_change)
        self.font_select.pack(side="left", padx=(0, 16))

        tk.Label(controls_line, text="Taille", bg=BG, fg=TEXT, font=("Arial", 15, "bold")).pack(side="left", padx=(0, 8))
        self.size_select = RoundSelect(controls_line, [str(s) for s in SIZE_OPTIONS], str(self.font_size), width=100, on_change=self._on_size_change)
        self.size_select.pack(side="left", padx=(0, 16))

        btn_w = 180
        self.btn_format = TextButton(controls_line, "Formatter", btn_w, self.on_format)
        self.btn_format.pack(side="left", padx=(0, 8))
        self.btn_copy = TextButton(controls_line, "Copier", btn_w, self.on_copy)
        self.btn_copy.pack(side="left", padx=(0, 8))
        self.btn_download = TextButton(controls_line, "Télécharger", btn_w, self.on_download_click)
        self.btn_download.pack(side="left", padx=(0, 8))

        self.btn_theme = TextButton(controls_line, "", btn_w, self.toggle_input_theme, accent=True)
        self.btn_theme.pack(side="left", padx=(0, 10))
        self._refresh_theme_button()

        self.font_label = tk.Label(controls_line, text="Police active: Arial", bg=BG, fg=MUTED, font=("Arial", 11, "bold"))
        self.font_label.pack(side="left", padx=(6, 0))

        self.split = tk.PanedWindow(
            main,
            orient=tk.HORIZONTAL,
            sashwidth=8,
            sashcursor="sb_h_double_arrow",
            showhandle=False,
            bg=BG,
            bd=0,
            relief="flat",
        )
        self.split.pack(fill="both", expand=True)

        left_card = RoundedCard(self.split, radius=14, pad=12)
        right_card = RoundedCard(self.split, radius=14, pad=12)
        self.split.add(left_card, minsize=350)
        self.split.add(right_card, minsize=350)

        tk.Label(left_card.inner, text="Texte source", bg=CARD, fg=TEXT, font=("Arial", 17, "bold")).pack(anchor="w", pady=(0, 8))
        self.input_text = tk.Text(
            left_card.inner,
            wrap="word",
            relief="flat",
            bd=0,
            highlightthickness=0,
            padx=10,
            pady=8,
            undo=True,
            insertbackground="#FFFFFF",
            cursor="xterm",
        )
        self.input_text.pack(fill="both", expand=True)

        tk.Label(right_card.inner, text="Texte formaté", bg=CARD, fg=TEXT, font=("Arial", 17, "bold")).pack(anchor="w", pady=(0, 8))
        self.output_text = tk.Text(
            right_card.inner,
            wrap="word",
            relief="flat",
            bd=0,
            highlightthickness=0,
            padx=10,
            pady=8,
            state="disabled",
            cursor="xterm",
        )
        self.output_text.pack(fill="both", expand=True)

        self.metrics = tk.Label(right_card.inner, text="FRE: - | Mots communs: -", bg=CARD, fg=MUTED, font=("Arial", 11, "bold"))
        self.metrics.pack(anchor="w", pady=(8, 0))

    def _preferred_font(self) -> tuple[str, bool]:
        families = set(tkfont.families())
        if self.font_name in families:
            return self.font_name, False

        if self.font_name == "OpenDyslexic":
            for alt in ("OpenDyslexic", "OpenDyslexic3", "Open Dyslexic", "OpenDyslexicAlta"):
                if alt in families:
                    return alt, False

        if self.font_name == "Avenir Next":
            for alt in ("Avenir Next", "Avenir Next Regular", "Avenir"):
                if alt in families:
                    return alt, False

        if "Arial" in families:
            return "Arial", True
        return "TkDefaultFont", True

    def _render_markdown(self, markdown: str, family: str, size: int) -> None:
        self.output_text.configure(state="normal")
        self.output_text.delete("1.0", "end")

        base_font = tkfont.Font(family=family, size=size, weight="normal")
        bold_font = tkfont.Font(family=family, size=size, weight="bold")
        self.output_text.tag_configure("base", font=base_font, spacing2=int(size * 0.5))
        self.output_text.tag_configure("bold", font=bold_font, spacing2=int(size * 0.5))

        for txt, is_bold in parse_markdown_segments(markdown):
            self.output_text.insert("end", txt, ("bold" if is_bold else "base",))

        self.output_text.configure(state="disabled")

    def _apply_input_theme(self) -> None:
        if self.input_dark_mode:
            bg = INPUT_DARK
            fg = INPUT_DARK_TEXT
        else:
            bg = INPUT_LIGHT
            fg = INPUT_LIGHT_TEXT

        self.input_text.configure(bg=bg, fg=fg, insertbackground=fg)
        self.output_text.configure(bg=bg, fg=fg, insertbackground=fg)

        if self.last_result:
            self.on_style_change()

    def _refresh_theme_button(self) -> None:
        label = "Mode saisie: sombre" if self.input_dark_mode else "Mode saisie: clair"
        self.btn_theme.set_text(label)

    def toggle_input_theme(self) -> None:
        self.input_dark_mode = not self.input_dark_mode
        self._refresh_theme_button()
        self._apply_input_theme()

    def _update_font_label(self, fallback: bool, active: str) -> None:
        if not self.font_label:
            return
        if fallback:
            self.font_label.config(text=f"Police active: {active} (fallback)")
        else:
            self.font_label.config(text=f"Police active: {active}")

    def _on_font_change(self, value: str) -> None:
        self.font_name = value
        self.on_style_change()

    def _on_size_change(self, value: str) -> None:
        try:
            self.font_size = int(value)
        except Exception:
            self.font_size = 13
        self.on_style_change()

    def on_style_change(self) -> None:
        if not self.last_result:
            active, fallback = self._preferred_font()
            self._update_font_label(fallback, active)
            return

        active, fallback = self._preferred_font()
        self._update_font_label(fallback, active)
        self._render_markdown(self.last_result.markdown_text, active, self.font_size)

    def on_format(self) -> None:
        source = self.input_text.get("1.0", "end").strip()
        if not source:
            messagebox.showwarning("Texte manquant", "Veuillez coller un texte à adapter.")
            return

        self.last_result = gaussReformulator(source)

        active, fallback = self._preferred_font()
        self._update_font_label(fallback, active)
        self._render_markdown(self.last_result.markdown_text, active, self.font_size)

        fre = self.last_result.analysis.flesch_score
        zipf = self.last_result.analysis.zipf_common_ratio * 100.0
        self.metrics.config(text=f"FRE: {fre:.1f} | Mots communs: {zipf:.1f}%")

    def on_copy(self) -> None:
        if not self.last_result:
            messagebox.showinfo("Aucun output", "Formatez un texte avant de copier.")
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(self.last_result.markdown_text)
        messagebox.showinfo("Copie", "Texte copié.")

    def on_download_click(self) -> None:
        if not self.last_result:
            messagebox.showinfo("Aucun output", "Formatez un texte avant de télécharger.")
            return

        if self.download_menu is None:
            self.download_menu = tk.Menu(self.root, tearoff=0)
            self.download_menu.add_command(label="Word (.docx)", command=self.export_docx)
            self.download_menu.add_command(label="PDF (.pdf)", command=self.export_pdf)
            self.download_menu.add_command(label="HTML (.html)", command=self.export_html)

        x = self.btn_download.winfo_rootx()
        y = self.btn_download.winfo_rooty() + self.btn_download.winfo_height() + 2
        self.download_menu.tk_popup(x, y)

    def export_html(self) -> None:
        if not self.last_result:
            return
        path = filedialog.asksaveasfilename(
            title="Exporter en HTML",
            defaultextension=".html",
            filetypes=[("HTML", "*.html")],
            initialfile="rumba_reader_output.html",
        )
        if not path:
            return

        active, _ = self._preferred_font()
        style = OutputStyle(font_family=active, font_size=self.font_size, line_height=1.5)
        doc = render_html_document(self.last_result, style)
        Path(path).write_text(doc, encoding="utf-8")
        messagebox.showinfo("Téléchargement", "HTML exporté.")

    def export_docx(self) -> None:
        if not self.last_result:
            return
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
        except Exception:
            messagebox.showerror("Module manquant", "Installez python-docx: pip install python-docx")
            return

        path = filedialog.asksaveasfilename(
            title="Exporter en Word",
            defaultextension=".docx",
            filetypes=[("Word", "*.docx")],
            initialfile="rumba_reader_output.docx",
        )
        if not path:
            return

        active, _ = self._preferred_font()
        lines = split_segments_by_line(parse_markdown_segments(self.last_result.markdown_text))

        doc = Document()
        rgb = RGBColor(255, 255, 255) if self.input_dark_mode else RGBColor(107, 114, 128)
        for line in lines:
            p = doc.add_paragraph()
            for txt, is_bold in line:
                run = p.add_run(txt)
                run.bold = is_bold
                run.font.name = active
                run.font.size = Pt(self.font_size)
                run.font.color.rgb = rgb

        doc.save(path)
        messagebox.showinfo("Téléchargement", "Word exporté.")

    def export_pdf(self) -> None:
        if not self.last_result:
            return
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib import colors
            from reportlab.platypus import Paragraph, SimpleDocTemplate
        except Exception:
            messagebox.showerror("Module manquant", "Installez reportlab: pip install reportlab")
            return

        path = filedialog.asksaveasfilename(
            title="Exporter en PDF",
            defaultextension=".pdf",
            filetypes=[("PDF", "*.pdf")],
            initialfile="rumba_reader_output.pdf",
        )
        if not path:
            return

        txt_color = colors.white if self.input_dark_mode else colors.HexColor("#6B7280")
        style = ParagraphStyle(
            name="RumbaPDF",
            fontName="Helvetica",
            fontSize=self.font_size,
            leading=self.font_size * 1.5,
            textColor=txt_color,
        )

        parts: List[str] = []
        for txt, is_bold in parse_markdown_segments(self.last_result.markdown_text):
            safe = xml_escape(txt).replace("\n", "<br/>")
            parts.append(f"<b>{safe}</b>" if is_bold else safe)

        doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
        doc.build([Paragraph("".join(parts), style)])
        messagebox.showinfo("Téléchargement", "PDF exporté.")


if __name__ == "__main__":
    root = tk.Tk()
    app = RumbaReaderApp(root)
    root.mainloop()
